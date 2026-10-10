"""Build the editable research presentation from versioned project results.
Run: python -m pip install python-pptx pillow; python ppt/build_presentation.py
All slide graphics are native PowerPoint shapes/text; no scientific result is AI-generated.
"""
from pathlib import Path
import csv, gzip, math, collections
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'ppt'; HIST=ROOT/'results/historical'
R=Presentation(); R.slide_width=Inches(13.333); R.slide_height=Inches(7.5)
NAVY='183249'; TEAL='168C91'; BLUE='527FA7'; ORANGE='CE8151'; GRAY='70818D'; LIGHT='E9F0F3'; WHITE='FFFFFF'; INK='243D4E'; RED='AC5960'
FONT='Microsoft YaHei'; previews=[]; slides=[]
fontpath='/mnt/c/Windows/Fonts/msyh.ttc'; boldpath='/mnt/c/Windows/Fonts/msyhbd.ttc'
SCALE=120

def data(folder,file):
 p=HIST/folder/file
 with (gzip.open(p,'rt',encoding='utf-8-sig') if p.suffix=='.gz' else p.open(encoding='utf-8-sig')) as f:return list(csv.DictReader(f))
def rgb(c):return RGBColor.from_string(c)
def box(x,y,w,h,c=LIGHT,outline=None,oval=False):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL if oval else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background();sh._element.spPr.append(OxmlElement('a:effectLst'))
 if outline:sh.line.color.rgb=rgb(outline)
 xy=(int(x*SCALE),int(y*SCALE),int((x+w)*SCALE),int((y+h)*SCALE))
 (d.ellipse if oval else d.rectangle)(xy,fill='#'+c,outline='#'+outline if outline else None)
 return sh

def text(t,x,y,w,h=0.45,size=20,color=INK,bold=False,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=Pt(0);tf.margin_top=tf.margin_bottom=Pt(0)
 for i,line in enumerate(str(t).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(3)
  for run in p.runs:
   ea=OxmlElement('a:ea');ea.set('typeface',FONT);run._r.get_or_add_rPr().append(ea)
 f=ImageFont.truetype(boldpath if bold else fontpath,int(size*SCALE/72))
 # Preview mirrors the native shape layout; Office/WPS rendering is checked separately if available.
 yy=y*SCALE
 for para in str(t).split('\n'):
  line=''
  for ch in para:
   if d.textlength(line+ch,font=f)>w*SCALE and line:
    d.text((x*SCALE,yy),line,font=f,fill='#'+color);yy+=size*SCALE/72*1.35;line=''
   line+=ch
  d.text((x*SCALE,yy),line,font=f,fill='#'+color);yy+=size*SCALE/72*1.35
 return sh

def line(x1,y1,x2,y2,color=GRAY,width=1.3):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);sh._element.spPr.append(OxmlElement('a:effectLst'))
 d.line((x1*SCALE,y1*SCALE,x2*SCALE,y2*SCALE),fill='#'+color,width=max(1,int(width*SCALE/72)))
def dot(x,y,color=TEAL,r=.055):box(x-r,y-r,r*2,r*2,color,oval=True)
def new(title,section,source,subtitle=''):
 global s,d,im
 s=R.slides.add_slide(R.slide_layouts[6]);im=Image.new('RGB',(1600,900),'white');d=ImageDraw.Draw(im)
 slides.append((s,im))
 box(0,0,13.333,.08,TEAL);text(section.upper(),.5,.22,12,.3,11,TEAL,True);text(title,.5,.72,12.35,.65,28,NAVY,True)
 if subtitle:text(subtitle,.52,1.43,12.2,.45,14,GRAY)
 line(.5,6.99,12.8,6.99,LIGHT);text(source,.52,7.08,11.6,.27,8,GRAY);text(f'{len(slides):02d}',12.32,7.06,.5,.28,10,GRAY)
 s.notes_slide.notes_text_frame.text='来源与解释：\n'+source+'\n'+subtitle+'\n图表均为原生可编辑形状；未显示的误差区间不代表不存在不确定性。'
 return s

def note(t):s.notes_slide.notes_text_frame.text+='\n'+t

def takeaway(t,color=TEAL):box(.5,6.35,12.3,.46,'EAF4F3');text(t,.68,6.40,11.9,.37,15,color,True)
def card(x,y,w,title,body,color=TEAL):
 box(x,y,w,2.25,'F3F6F8');box(x,y,.055,2.25,color);text(title,x+.18,y+.2,w-.36,.6,21,color,True);text(body,x+.18,y+.91,w-.36,1.1,17)
def flow(items,y=2.5):
 w=11.8/len(items)-.27
 for i,(a,b) in enumerate(items):
  x=.65+i*(w+.32);box(x,y,w,1.75,'F0F5F7');text(a,x+.15,y+.16,w-.3,.6,19,TEAL,True);text(b,x+.15,y+.85,w-.3,.75,15)
  if i<len(items)-1:text('→',x+w+.02,y+.55,.3,.5,20,GRAY)
def bars(labels,values,x=.7,y=2.25,w=8.0,h=3.6,maxv=.6,colors=None):
 labelw=2.55; left=x+labelw;plotw=w-labelw-.6
 for tick in range(5):
  v=maxv*tick/4;xx=left+plotw*v/maxv;line(xx,y-.13,xx,y+h,LIGHT,.8);text(f'{v:.2f}',xx-.18,y+h+.08,.6,.3,10,GRAY)
 for i,(lab,v) in enumerate(zip(labels,values)):
  yy=y+i*h/len(values);text(lab,x,yy+.06,labelw-.12,.5,15);box(left,yy+.08,plotw*v/maxv,min(.42,h/len(values)*.58),(colors or [TEAL]*len(labels))[i]);text(f'{v:.4f}',left+plotw*v/maxv+.1,yy+.08,.8,.35,13,INK,True)
 text('NDCG@10（越高越好）',left,y+h+.45,4,.35,12,GRAY)
def table(headers,rows,widths,y=2.05,rowh=.62,fs=16):
 x=.6
 for h,w in zip(headers,widths):box(x,y,w,.55,NAVY);text(h,x+.12,y+.12,w-.24,.35,fs-1,WHITE,True);x+=w
 for i,row in enumerate(rows):
  x=.6
  for j,(v,w) in enumerate(zip(row,widths)):
   box(x,y+.55+i*rowh,w,rowh,'F0F5F7' if i%2==0 else WHITE);text(v,x+.12,y+.65+i*rowh,w-.24,rowh-.1,fs,TEAL if j==0 else INK,j==0);x+=w

def forest(labels,means,lows,highs,xmin=-.25,xmax=.05,y=2.55):
 left=4.0;right=11.9
 for v in [xmin, (xmin+xmax)/2,0,xmax]:
  xx=left+(v-xmin)/(xmax-xmin)*(right-left);line(xx,y-.35,xx,y+2.55,GRAY if v==0 else LIGHT,1);text(f'{v:+.2f}',xx-.25,y+2.67,.7,.3,12,GRAY)
 for i,(lab,m,lo,hi) in enumerate(zip(labels,means,lows,highs)):
  yy=y+i*1.1;text(lab,.65,yy-.15,3.2,.7,18)
  xx=lambda v:left+(v-xmin)/(xmax-xmin)*(right-left)
  line(xx(lo),yy+.07,xx(hi),yy+.07,TEAL,3);dot(xx(m),yy+.07,TEAL,.075);text(f'{m:+.4f}  [{lo:+.4f}, {hi:+.4f}]',4.4,yy+.3,6,.4,15)

def curve(series,epochs,x=.8,y=2.3,w=7.8,h=3.0,lo=.20,hi=.32):
 for v in [lo,(lo+hi)/2,hi]:
  yy=y+h-(v-lo)/(hi-lo)*h;line(x+.5,yy,x+w,yy,LIGHT);text(f'{v:.2f}',x,yy-.13,.5,.3,11,GRAY)
 for label,vals,c in series:
  pts=[(x+.65+i*(w-.9)/(len(epochs)-1),y+h-(v-lo)/(hi-lo)*h) for i,v in enumerate(vals)]
  for a,b in zip(pts,pts[1:]):line(*a,*b,c,2.3)
  for xx,yy in pts:dot(xx,yy,c,.045)
 for i,e in enumerate(epochs):text(str(e),x+.48+i*(w-.9)/(len(epochs)-1),y+h+.12,.55,.3,11,GRAY)
 text('Epoch（检查点等距显示）',x+2.7,y+h+.48,4,.4,12,GRAY)

# 01
new('泛癌选择性基因依赖预测\n与 ccRCC 候选优先排序','RESEARCH PRESENTATION','PanCancer-Dependency | 结果截至 2026-09-24；汇报版本 2026-10-10')
text('跨癌系泛化 · 跨平台复现 · 模型公平比较 · 患者证据边界',.6,2.22,12,.6,21,GRAY)
flow([('表达谱','细胞系 / 患者'),('选择性依赖','排除普遍必需基因'),('外部检验','CRISPR / RNAi'),('候选证据','功能与正常肾风险')],3.35)
text('核心发现：表达含有可迁移信号；复杂模型尚未稳定超过 PCR-ridge。',.65,5.75,12,.5,22,TEAL,True)
#02
new('研究问题：从分子特征到可验证的功能假设','01 / RATIONALE','README.md：研究目的、项目最终能够支持的表述')
card(.65,2.1,3.85,'预测什么？','哪些基因对特定癌模型\n比平均水平更重要？')
card(4.75,2.1,3.85,'能否迁移？','在未见癌系和另一扰动\n平台上，排序是否仍有效？')
card(8.85,2.1,3.85,'证据到哪一步？','患者候选是否稳定？\n正常肾是否存在功能风险？',ORANGE)
text('功能依赖：基因受扰动后，细胞生长或存活受到影响。',.8,5.0,11,.6,21)
takeaway('研究目标是缩小功能实验范围；当前尚无患者匹配的基因扰动真值。')
#03
new('公开数据构成分层证据链','01 / DATA LANDSCAPE','README.md：数据和评价方法；results/historical/final_evidence_synthesis_v2/')
table(['数据层','规模 / 角色','能够回答的问题'],[
 ('DepMap 24Q4','873 模型；15,835 共有基因','表达能否预测细胞系 CRISPR 依赖？'),
 ('跨癌系评价','805 模型；19 癌系；780 患者组','能否泛化到训练中未见的癌系？'),
 ('Sanger / DRIVE','外部 CRISPR / 独立 RNAi 技术','依赖方向能否跨扰动平台复现？'),
 ('TCGA-KIRC','156 Train / 53 Validation / 51 Test','患者候选频率与表达方向是否稳定？'),
 ('HPA / GTEx / 文献','正常组织暴露与直接扰动证据','候选的支持、冲突与正常肾风险？')],[2.35,4.0,5.75],rowh=.68,fs=16)
takeaway('数据层级决定结论层级：表达稳定性不能替代功能验证。')
#04
new('分析框架：模型比较与患者证据分别审查','02 / STUDY DESIGN','README.md；configs/deepdep_recovered_comparison_protocol_20260921.json')
text('A  模型评价',.65,2,3,.5,21,TEAL,True)
flow([('表达 → 依赖','训练折内预处理'),('癌系整组留出','患者重叠清除'),('强基线比较','PCR / EN / RRR / Deep'),('Sanger 复查','按冻结门槛推进')],2.7)
text('B  历史患者分析',.65,4.8,3,.4,21,ORANGE,True)
text('冻结核岭模型 → TCGA 映射 → 前 20 候选冻结 → Locked Test → 分层证据整合',.65,5.4,12,.6,20)
takeaway('患者候选来自历史冻结核模型，不能改称最新 PCR 或 DeepDEP 的候选。')
#05
new('依赖残差：区分普遍必需与相对选择性','02 / TARGET DEFINITION','outputs/selective_dependency_lolo_v1/；README.md：阶段二','概念示意；下方矩阵不代表任何实际基因或样本。')
text('r(i,j) = y(i,j) − mean_train(j)',.85,2.1,7,.7,30,TEAL,True)
text('每个外层训练折单独计算基因均值；更负的残差表示更强的相对依赖。',.85,3.0,11.5,.6,20)
for j,title in enumerate(['普遍强依赖','谱系选择性','低依赖']):
 x=1.15+j*4;text(title,x,4.1,3,.4,18,NAVY,True)
 for k,c in enumerate(([TEAL]*6,[LIGHT,LIGHT,TEAL,TEAL,LIGHT,LIGHT],[LIGHT]*6)[j]):box(x+k*.42,4.85,.35,.65,c)
text('主评价排除 release-wide common-essential；该注释未在每折内部重新估计。',.85,5.7,11.5,.4,16,ORANGE)
takeaway('残差反映相对依赖；残差低不自动意味着患者特异性或安全治疗窗。')
#06
new('整癌系留出与嵌套选择限制信息泄漏','02 / EVALUATION','configs/deepdep_recovered_comparison_protocol_20260921.json；deepdep_recovered_comparison_v1/run.json')
for i in range(19):box(.8+i*.61,2.45,.49,.75,ORANGE if i==8 else TEAL)
text('18 个癌系：外层训练',.8,2.0,7,.4,20,TEAL,True);text('1 个癌系：外层评估',8.0,3.45,4.7,.4,20,ORANGE,True)
line(6,3.3,6,4.0,GRAY);box(.8,4.1,7.3,1.55,LIGHT);text('训练内部：5 个癌系折选择超参数 / epoch',1,4.28,6.9,.5,20,NAVY,True);text('训练内 PCA、标准化与残差均值；排除患者重叠',1,5,6.9,.4,16)
box(8.55,4.1,3.9,1.55,'F8EEE5');text('最终比较',8.75,4.3,3.5,.5,20,ORANGE,True);text('19 癌系等权 + 配对 bootstrap',8.75,5,3.5,.6,16)
takeaway('最新模型比较属于开发期分析；不能称为未触碰的外部 Locked Test。')
#07
new('表达携带稳定的选择性依赖排序信息','03 / EXPRESSION SIGNAL','outputs/selective_dependency_lolo_v1/；outputs/selective_dependency_analysis_v1/')
bars(['训练选择性先验','表达残差模型'],[.2248,.4178],h=2.7,maxv=.5)
text('+0.1930',9,2.3,3,.7,34,TEAL,True);text('癌系等权 ΔNDCG',9,3.08,3,.5,18)
text('19 / 19',9,4.1,3,.7,32,TEAL,True);text('癌系增益方向为正',9,4.9,3,.4,18)
takeaway('支持表达具有可迁移的排序信息；12 个 ccRCC 模型不足以证明额外亚型优势。')
note('该页为历史广基因空间残差评价，不可与后续 911 靶点比较表混排。')
#08
new('跨平台复现为正，但具体基因名单准确率有限','03 / EXTERNAL CRISPR','outputs/sanger_selective_dependency_v1/；README.md：阶段三','历史冻结流程：直接表达 66 模型；映射表达 17 模型。')
forest(['直接表达','映射表达'],[.1065,.0576],[.0776,-.0099],[.1360,.1285],-.04,.18,y=2.6)
text('ΔNDCG 相对训练标准化先验；95% bootstrap 区间',3.7,5.7,8,.4,15,GRAY)
takeaway('直接表达绝对 NDCG = 0.1670，真实 Top-10 重叠 = 0.0303；增益不等于高准确率。')
#09
new('公平比较必须固定输入、靶点与评价口径','04 / MODEL BENCHMARKS','results/historical/advanced_model_benchmark_v1/；sanger_advanced_benchmark_v1/')
card(.65,2.1,3.85,'相同输入','同一表达矩阵\n同一模型 / 患者划分')
card(4.75,2.1,3.85,'相同目标','1,204 个映射靶点\n主评价 911 个非普遍必需基因')
card(8.85,2.1,3.85,'相同主终点','癌系等权 NDCG@10\n训练内调参 + 配对区间')
text('历史广基因空间、受限靶点空间、恢复表达后的比较：分别报告。',.8,5.2,11.9,.6,21,ORANGE,True)
takeaway('不能把不同基因空间、不同外部校准或不同聚合权重的分数混成一个排行榜。')
#10
new('增加模型复杂度没有带来稳定的主终点增益','04 / RESTRICTED TARGET BENCHMARK','results/historical/advanced_model_benchmark_v1/overall_summary.csv；bootstrap.csv','首次受限靶点比较：911 个非 common-essential 基因；19 癌系等权。')
adv=data('advanced_model_benchmark_v1','overall_summary.csv');ad={r['method']:float(r['ndcg_at_10']) for r in adv if r['estimand']=='lineage_equal'}
keys=['expression_pcr_ridge','multitask_reduced_rank_ridge','per_target_elastic_net_shared','per_target_elastic_net_targetwise','exp_deepdep_adapted']
bars(['PCR-ridge','多任务低秩 RRR','Elastic Net：共享 α','Elastic Net：逐靶点 α','Exp-DeepDEP 适配'],[ad[k] for k in keys],h=3.2,maxv=.55,colors=[TEAL,BLUE,GRAY,GRAY,ORANGE])
text('RRR − PCR',9.1,2.35,3,.4,20,TEAL,True);text('+0.0009',9.1,3.0,3,.6,29,TEAL,True);text('95% CI\n[−0.0043, +0.0060]',9.1,3.85,3.3,1,18)
takeaway('RRR 的点估计接近 PCR；区间跨 0，不支持主排序指标稳定获益。')
#11
new('数值审计排除了部分优化精度解释','04 / ELASTIC NET AUDIT','results/historical/elastic_net_numerical_audit_v1/；elastic_net_upper_grid_v1/')
table(['实验','NDCG@10','解释'],[('PCR-ridge','0.4585','同一受限靶点基准'),('原 Elastic Net','0.4428','原共享 α 拟合'),('严格数值精化','0.4405','精化后未追平 PCR'),('扩大 α 上界','0.4405','19/19 癌系仍选择 α = 0.1')],[3.2,2.1,6.8],rowh=.7,fs=18)
text('上界网格 0.1 / 0.2 / 0.5 / 1.0；480,396 次靶点拟合通过数值认证。',.8,5.7,11.7,.4,16,GRAY)
takeaway('结论仅覆盖已审计配置；更小 α 的完整选择过程及其他 l1_ratio 尚未穷尽。')
#12
new('恢复表达特征后重新配对比较 PCR 与 DeepDEP','05 / INPUT RECOVERY','results/historical/deepdep_expression_recovery_v1/；deepdep_recovered_comparison_v1/')
for y,label,obs,missing in [(2.65,'修复前',4678,1338),(4.15,'修复后',5384,632)]:
 text(label,.8,y,.95,.5,19,NAVY,True);box(2,y,8.7*obs/6016,.65,TEAL);box(2+8.7*obs/6016,y,8.7*missing/6016,.65,LIGHT);text(f'{obs:,} 实测 / {missing:,} 填补',2,y+.85,8,.4,19)
text('+706',11,3,2,.7,30,TEAL,True);text('恢复特征',11,3.85,2,.5,17)
takeaway('两类模型共享恢复后的输入；仍有 632 个位置需要填补，不能称完整原论文复现。')
#13
new('两条模型路径使用同一恢复表达输入','05 / MODEL ARCHITECTURE','configs/deepdep_recovered_comparison_protocol_20260921.json','Exp-DeepDEP 为表达适配实现；架构示意不表示各层尺寸比例。')
text('PCR-ridge',.7,2.0,3,.4,21,TEAL,True)
for i,(a,b) in enumerate([('表达','6,016 位置'),('训练内 PCA','rank 32–256'),('Ridge','α 100–100,000'),('依赖残差','多靶点输出')]):
 x=.7+i*3.12;box(x,2.65,2.7,.95,LIGHT);text(a,x+.12,2.77,2.45,.35,19,TEAL,True);text(b,x+.12,3.17,2.45,.3,13)
 if i<3:text('→',x+2.76,2.88,.3,.4,20)
text('Exp-DeepDEP',.7,4.02,3,.4,21,ORANGE,True)
for yy,a,b in [(4.58,'表达编码器','6016 → 500 → 200 → 50'),(5.37,'基因指纹编码器','3115 → 1000 → 100 → 50')]:
 box(.7,yy,4.7,.66,LIGHT);text(a,.85,yy+.06,1.8,.35,15,TEAL,True);text(b,2.67,yy+.12,2.6,.35,13)
 line(5.4,yy+.33,6.0,yy+.33,GRAY);line(6.0,yy+.33,6.0,5.3,GRAY)
line(6.0,5.3,6.45,5.3,GRAY);text('→',6.12,4.99,.4,.4,20)
box(6.55,4.88,3,.96,LIGHT);text('拼接 → 依赖预测',6.7,5.0,2.75,.4,18,TEAL,True);text('100 → 250 → 250 → 1',6.7,5.47,2.75,.3,14)
text('500 cell–target pairs / batch',9.9,4.94,2.7,.8,17,ORANGE,True)
takeaway('DeepDEP 使用既有非 KIRC 肿瘤预训练；主分析为残差目标，绝对依赖为次分析。')
#14
new('训练成本来自嵌套折、随机种子与成对批次','05 / COMPUTATIONAL BUDGET','results/historical/deepdep_recovered_comparison_v1/training.csv.gz；run.json')
flow([('19','外层留出癌系'),('5 × 3 × 2','内层折 × 种子 × 目标'),('570','内层拟合，各 100 epoch'),('+ 114','外层最终重拟合')],2.1)
text('684',.9,4.55,3,.7,36,TEAL,True);text('完成拟合数',.9,5.36,3,.4,19)
text('61,710',5.0,4.55,3.6,.7,36,TEAL,True);text('累计 epoch 数',5.0,5.36,3.6,.4,19)
text('≈ 49.9 h',9.2,4.55,3.4,.7,34,ORANGE,True);text('记录的 epoch 耗时之和',9.2,5.36,3.6,.45,17)
takeaway('49.9 h 不是总墙钟时间，也不是纯 GPU 利用时长；不能据此比较所有模型总成本。')
#15
new('恢复表达后，PCR 仍明显优于两种 DeepDEP 目标','06 / PRIMARY RESULT','results/historical/deepdep_recovered_comparison_v1/overall_summary.csv','19 癌系等权；805 模型；911 个非 common-essential 靶点。')
latest=data('deepdep_recovered_comparison_v1','overall_summary.csv');ld={r['method']:float(r['ndcg_at_10']) for r in latest if r['estimand']=='lineage_equal'}
methods=['expression_pcr_ridge','exp_deepdep_residual_primary','exp_deepdep_absolute_secondary']
bars(['PCR-ridge','DeepDEP 残差：主分析','DeepDEP 绝对值：次分析'],[ld[k] for k in methods],h=3,maxv=.55,colors=[TEAL,ORANGE,BLUE])
text('0.4628',9.2,2.4,3,.7,33,TEAL,True);text('PCR NDCG@10',9.2,3.2,3,.5,18)
text('0.4066',9.2,4.2,3,.7,30,TEAL,True);text('PCR 选择性命中率@10',9.2,5,3.2,.4,15)
takeaway('在当前数据、训练协议和评价目标下，没有证据支持用 DeepDEP 替换 PCR。')
#16
new('配对区间与逐癌系方向一致支持 PCR','06 / PAIRED UNCERTAINTY','results/historical/deepdep_recovered_comparison_v1/bootstrap.csv；lineage_summary.csv','差值 = DeepDEP − PCR；按癌系等权 bootstrap，n = 19。')
br=data('deepdep_recovered_comparison_v1','bootstrap.csv');sel=[]
for m in methods[1:]:sel.append(next(r for r in br if r['comparison']==m+'_vs_expression_pcr_ridge' and r['metric']=='ndcg_at_10' and r['estimand']=='lineage_equal'))
forest(['残差主分析 − PCR','绝对值次分析 − PCR'],*[ [float(r[k]) for r in sel] for k in ['mean','ci_low','ci_high']])
takeaway('两种 DeepDEP 目标均在 19/19 个癌系低于 PCR；该结果不能外推为所有深度学习无效。')
#17
new('癌系异质性存在，但未出现 DeepDEP 优势癌系','06 / LINEAGE-LEVEL RESULTS','results/historical/deepdep_recovered_comparison_v1/lineage_summary.csv','每行一个留出癌系；颜色与数值均表示 NDCG@10，不显示独立误差区间。')
lin=data('deepdep_recovered_comparison_v1','lineage_summary.csv');lineages=sorted(set(r['heldout_lineage'] for r in lin));lookup={(r['heldout_lineage'],r['method']):float(r['ndcg_at_10']) for r in lin}
for panel,subset in enumerate([lineages[:10],lineages[10:]]):
 x=.7+panel*6.25
 text('癌系',x,2,2,.4,14,GRAY)
 for j,name in enumerate(['PCR','Deep 残差','Deep 绝对']):text(name,x+2.4+j*1.02,2,1.0,.35,12,GRAY)
 for i,l in enumerate(subset):
  y=2.52+i*.34;text(l.replace('Esophagus/Stomach','Esophagus').replace('Peripheral Nervous System','PNS').replace('Ovary/Fallopian Tube','Ovary').replace('Head and Neck','Head / Neck'),x,y,2.4,.3,13)
  for j,m in enumerate(methods):
   v=lookup[l,m];a=min(1,v/.65);c=''.join(f'{int(245+(target-245)*a):02X}' for target in (22,140,145));box(x+2.4+j*1.02,y,.92,.30,c);text(f'{v:.3f}',x+2.54+j*1.02,y+.02,.8,.26,11,WHITE if v>.38 else NAVY)
takeaway('Kidney：PCR 0.5661，Deep 残差 0.3238，Deep 绝对值 0.3070；Kidney 不等同于 ccRCC。')
#18
new('延长同类训练缺乏明确收益依据','06 / TRAINING DYNAMICS','results/historical/deepdep_recovered_comparison_v1/tuning.csv.gz','曲线是重叠训练任务的描述性汇总；不能当作独立样本推断，亦未据此重新选 epoch。')
tune=data('deepdep_recovered_comparison_v1','tuning.csv.gz');epochs=[1,2,5,10,20,40,60,80,100];groups=collections.defaultdict(list)
for r in tune:
 if r['method'] in methods[1:]:groups[r['method'],int(r['epoch'])].append(float(r['ndcg']))
vals=[[sum(groups[m,e])/len(groups[m,e]) for e in epochs] for m in methods[1:]]
curve([('残差',vals[0],ORANGE),('绝对值',vals[1],BLUE)],epochs)
text('内层平均 NDCG@10',1,1.92,5,.35,16,GRAY)
box(9.1,2.5,.25,.2,ORANGE);text('残差目标',9.55,2.4,2.7,.4,18)
box(9.1,3.2,.25,.2,BLUE);text('绝对值目标',9.55,3.1,2.7,.4,18)
text('性能未随 epoch\n持续上升',9.1,4.25,3.6,1,23,NAVY,True)
takeaway('当前证据支持停止追加同类训练；输入恢复与批次修复的各自因果贡献尚未分离。')
#19
new('受限靶点的 Sanger 外部比较同样未支持复杂模型','06 / EXTERNAL BASELINES','results/historical/sanger_advanced_benchmark_v1/；README.md：阶段十','历史高级模型外部比较：64 模型 / 13 癌系 / 871 共同评价基因；不是最新 DeepDEP 的外部结果。')
bars(['PCR-ridge','多任务低秩 RRR','Elastic Net','Exp-DeepDEP 适配','训练先验'],[.3351,.3270,.3029,.1665,.0725],h=3.2,maxv=.4,colors=[TEAL,BLUE,GRAY,ORANGE,GRAY])
text('RRR − PCR',9,2.4,3,.4,20,TEAL,True);text('−0.0082',9,3.1,3,.6,29,TEAL,True);text('95% CI\n[−0.0336, +0.0159]',9,4,3.4,1,18)
takeaway('最新恢复表达 DeepDEP 未通过内部推进门槛，因此没有新一轮 Sanger 验证。')
#20
new('患者域：队列稳定性与映射敏感性同时存在','07 / PATIENT TRANSFER','outputs/tcga_patient_transfer_v1/；results/historical/tcga_locked_test_v1/','本节均来自历史冻结核模型；患者无功能依赖标签。')
labels=['driver Kidney 中性化','driver non-Kidney 中性化','映射替代假设（开发）','映射替代假设（Test）'];v=[.9909,.9837,.3651,.3608]
# Generic horizontal plot with explicit overlap axis
for i,(lab,val) in enumerate(zip(labels,v)):
 yy=2.35+i*.78;text(lab,.8,yy,4.1,.5,17);box(5,yy+.02,6*val,.39,TEAL if i<2 else ORANGE);text(f'{val:.4f}',5+6*val+.12,yy,.9,.4,17)
text('患者 Top-10 名单平均重叠（0–1）',5,5.7,6,.4,16,GRAY)
takeaway('表达映射是重要不确定性：合理映射之间仅约 36% 的头部候选相同。')
#21
new('Locked Test 确认了稳定性，没有确认患者功能','07 / LOCKED TEST','results/historical/tcga_locked_test_v1/；冻结提交 73f9972')
flow([('冻结协议','模型 / 映射 / 候选 / 指标'),('一次正式访问','51 个患者'),('频率稳定性','前 20：ρ = 0.7984'),('配对表达方向','14 对肿瘤 / 正常')],2.15)
text('区间支持预设表达方向',.85,4.5,6,.45,20,TEAL,True);text('PAX8 · FERMT2 · CCND1',.85,5.1,6,.5,23)
text('方向相同，但区间跨 0',7.5,4.5,5,.45,20,ORANGE,True);text('HNF1B：尚未确认',7.5,5.1,5,.5,23)
takeaway('Test 已被正式访问；不能再充当新 PCR 或新 DeepDEP 的独立确认集。')
note('ρ=0.7984 指冻结前20候选 Validation–Test 出现频率 Spearman。')
#22
new('证据逐层收窄：计算候选不等同于治疗靶点','08 / CANDIDATE EVIDENCE','results/historical/final_evidence_synthesis_v2/candidate_evidence_matrix.csv；ccrcc_primary_functional_audit_v1/','患者来源资源结论限于 2026 年 9 月冻结审计范围，并非声称全球不存在相关数据。')
for i,(n,label,width,c) in enumerate([(20,'冻结计算候选',10,TEAL),(10,'DRIVE 文库覆盖',8,BLUE),(3,'ccRCC RNAi 绝对残差区间 < 0',6.1,BLUE),(1,'相对其他 Kidney 的 RNAi 区间 < 0',4.3,ORANGE)]):
 yy=2.05+i*.88;box(1.5+(10-width)/2,yy,width,.66,c);text(str(n),1.65+(10-width)/2,yy+.08,.75,.4,23,WHITE,True);text(label,2.45+(10-width)/2,yy+.15,width-1.1,.4,15,WHITE)
text('0 个候选获得该审计范围内的患者来源功能真值确认',1.65,5.75,10.5,.4,19,RED,True)
takeaway('DRIVE 是独立扰动技术，部分细胞系重叠；不能称为独立患者队列验证。')
#23
new('主要候选的功能支持与风险并存','08 / EVIDENCE MATRIX','results/historical/final_evidence_synthesis_v2/candidate_evidence_matrix.csv','强调证据层次；候选展示顺序用于讨论，未改变冻结 discovery rank。')
table(['候选','RNAi 依赖','相对 Kidney','直接文献','正常肾风险'],[
('PAX8','支持','支持','支持','表达 + 类器官'),('HNF1B','支持','区间跨 0','非同行评议','表达 + 类器官'),('FERMT2','支持','未建立','证据缺口','表达暴露'),('CCND1','区间跨 0','未建立','支持','治疗窗未知'),('YPEL5','未形成支持','未建立','反向证据','非优先抑制靶点')],[1.7,2.1,2.25,2.4,3.65],rowh=.66,fs=17)
takeaway('PAX8 证据最一致，但没有证明治疗窗；YPEL5 存在与优先抑制方向冲突的证据。')
#24
new('正交 RNAi 支持不能消除正常肾暴露问题','08 / FUNCTION VERSUS EXPOSURE','results/historical/final_evidence_synthesis_v2/candidate_evidence_matrix.csv','左：DRIVE ccRCC 平均依赖残差与 bootstrap 区间；右：GTEx 肾组织最高中位 TPM。')
ev=data('final_evidence_synthesis_v2','candidate_evidence_matrix.csv');evd={r['Gene']:r for r in ev}
for i,g in enumerate(['PAX8','HNF1B','FERMT2']):
 r=evd[g];yy=2.65+i*.94;text(g,.75,yy-.15,1.3,.4,20,TEAL,True)
 xx=lambda v:2.5+(v+1.3)/1.5*4.0
 lo,mu,hi=[float(r[k]) for k in ['drive_ccrcc_residual_ci_low','drive_ccrcc_residual_mean','drive_ccrcc_residual_ci_high']]
 line(xx(lo),yy,xx(hi),yy,TEAL,3);dot(xx(mu),yy,TEAL,.07)
 tpm=float(r['gtex_kidney_max_median_tpm']);box(8,yy-.18,3.1*tpm/400,.36,ORANGE);text(f'{tpm:.1f}',8+3.1*tpm/400+.13,yy-.2,1,.4,17)
for tick in [-1.2,-.6,0]:
 line(xx(tick),5.1,xx(tick),5.2,GRAY);text(f'{tick:g}',xx(tick)-.18,5.25,.7,.3,12)
for tick in [0,200,400]:
 xxx=8+3.1*tick/400;line(xxx,5.1,xxx,5.2,GRAY);text(str(tick),xxx-.1,5.25,.6,.3,12)
line(xx(0),2.1,xx(0),5.1,GRAY);text('更负 = 更强相对依赖',2.5,5.65,4.7,.4,15,GRAY);text('0–400 TPM；暴露而非毒性',8,5.65,4.7,.4,15,GRAY)
takeaway('正常肾类器官发育扰动提示风险；不能直接换算为成人肾毒性或临床安全窗。')
#25
new('与既有工作的区别在评价问题与证据审计','09 / POSITIONING','Shi et al. Nature Cancer (2024), doi:10.1038/s43018-024-00789-y；DeepDEP (2021), PMC8378822','文献差异不是性能优越性的证据；当前适配实现并非原论文逐项复现。')
table(['维度','TCGA-DEPMAP / DeepDEP','本项目'],[('核心目标','泛癌依赖预测与患者转化','选择性依赖排序与 ccRCC 候选审计'),('当前侧重','患者关联、转化资源与实验支持','整癌系留出、强基线、数值与输入审计'),('患者证据','各自设计的临床关联 / 实验验证','冻结 Test 的表达与频率稳定性'),('结论边界','不同数据与评价协议','不能用本项目分数宣称超过原论文')],[2.15,5.0,4.95],rowh=.8,fs=17)
takeaway('合理定位：可审查的预测与证据流程；创新性不能仅由模型复杂度或候选名称支持。')
note('Primary source: https://www.nature.com/articles/s43018-024-00789-y ; DeepDEP https://pmc.ncbi.nlm.nih.gov/articles/PMC8378822/ . Prior-work table is historical; the new model comparison does not replace historical patient predictions.')
#26
new('当前主要限制及下一阶段的决策门槛','10 / LIMITATIONS AND NEXT STEPS','README.md：现有不足和偏差来源、下一步工作')
card(.65,2.1,3.85,'方法层面','统一比较表与逐癌系图\n独立环境完整复跑\n保留 PCR 作为强基线')
card(4.75,2.1,3.85,'患者迁移','单独冻结新模型协议\n报告映射敏感性\n另寻未使用确认数据',BLUE)
card(8.85,2.1,3.85,'功能与安全','患者来源模型扰动\n正常肾平行扰动\n剂量与可比生长终点',ORANGE)
text('公开数据阶段的主要瓶颈：功能真值缺失、少量 ccRCC 模型、细胞系到患者的域差异。',.8,5.2,11.8,.7,20)
takeaway('若目标升级为治疗靶点结论，必须补充新的功能证据；继续同类训练不能填补这一缺口。')
#27
new('结论：表达可预测，复杂度不保证收益，患者证据仍有限','11 / TAKE-HOME MESSAGES','README.md；results/historical/deepdep_recovered_comparison_v1/；final_evidence_synthesis_v2/')
for i,(num,title,body,c) in enumerate([('01','方法结论','表达支持可迁移的细胞系选择性依赖排序。',TEAL),('02','模型结论','同数据严格比较未发现复杂模型稳定超过 PCR。',BLUE),('03','生物学结论','获得分层候选与风险线索，尚未证明患者功能依赖或治疗窗。',ORANGE)]):
 y=2.05+i*1.25;text(num,.75,y,1,.65,32,c,True);text(title,2,y,3,.45,21,c,True);text(body,2,y+.55,10.4,.6,21)
takeaway('项目成果是可追溯的候选优先排序与证据边界，不是已验证的临床靶点。')
#28
new('数据与结果来源','APPENDIX / REPRODUCIBILITY','所有本项目数值可追溯到仓库 README、冻结配置与版本化 CSV；图形为原生可编辑对象。')
refs=[('DepMap 24Q4','doi:10.25452/figshare.plus.27993248.v1'),('TCGA-DEPMAP','Shi et al. Building a translational cancer dependency map for The Cancer Genome Atlas.\nNature Cancer 5, 1176–1194 (2024). doi:10.1038/s43018-024-00789-y'),('DeepDEP','2021 原文：pmc.ncbi.nlm.nih.gov/articles/PMC8378822/\n官方实现：github.com/chenlab-gccri/DeepDEP'),('正式模型结果','results/historical/deepdep_recovered_comparison_v1/\nadvanced_model_benchmark_v1/；sanger_advanced_benchmark_v1/'),('候选与患者证据','results/historical/final_evidence_synthesis_v2/；tcga_locked_test_v1/\n完整溯源见各页备注和项目 README。')]
for i,(a,b) in enumerate(refs):
 y=2.0+i*.83;text(a,.7,y,2.4,.5,17,TEAL,True);text(b,3.05,y,9.5,.8,14)
# Metadata and save
R.core_properties.title='泛癌选择性基因依赖预测与 ccRCC 候选优先排序'
R.core_properties.subject='学术汇报；公开数据；严格基线比较与证据审计'
R.core_properties.author='PanCancer-Dependency';R.core_properties.keywords='ccRCC; PCR-ridge; DeepDEP; CRISPR; RNAi; editable'
path=OUT/'PanCancer_Dependency_学术汇报.pptx';R.save(path)
# Contact sheet checks the generator layout, not a claim of native Office rendering.
thumbs=[]
for j,(_,img) in enumerate(slides):
 t=img.resize((480,270));thumbs.append(t)
contact=Image.new('RGB',(480*4,290*math.ceil(len(thumbs)/4)), '#DCE5EA')
for i,t in enumerate(thumbs):contact.paste(t,((i%4)*480,(i//4)*290))
contact.save(OUT/'slide_overview.png')
print(f'【完成】{len(slides)} 页｜全部正文与图形可编辑｜{path}')
